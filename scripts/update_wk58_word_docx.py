#!/usr/bin/env python3
"""Write the updated Week 8 Project Plan Word file.

Adds the Gibran epigraph in front matter and a usage/tracking appendix
with the inclusion statements. Preserves existing Word comments.
Does not rewrite Alignment, Gap, method, or spoken interview text.
"""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph

SRC = Path("/workspace/downloads/MW_Wk5_to_wk8_Edit_6Sep26.docx")
OUT = Path("/workspace/downloads/MW_Wk5_to_wk8_Edit_6Sep26.docx")
ART = Path("/opt/cursor/artifacts/MW_Wk5_to_wk8_Edit_6Sep26.docx")
ALIAS = Path("/workspace/downloads/MW_Wk5_to_wk8_Edit_6Sep26_updated.docx")
ART_ALIAS = Path("/opt/cursor/artifacts/MW_Wk5_to_wk8_Edit_6Sep26_updated.docx")

TRACK_HEADING = "Usage and Tracking Update (6 September 2026)"
EPIGRAPH_HEADING = "Epigraph"


def set_run_font(run, *, size=12, bold=None, italic=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def insert_paragraph_after(paragraph, text: str = "", *, bold=False, italic=False, center=False) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        run = new_para.add_run(text)
        set_run_font(run, bold=bold, italic=italic)
    new_para.paragraph_format.line_spacing = 2.0
    if center:
        new_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return new_para


def insert_block(after: Paragraph, lines: list[dict]) -> Paragraph:
    cur = after
    for spec in lines:
        cur = insert_paragraph_after(
            cur,
            spec.get("text", ""),
            bold=spec.get("bold", False),
            italic=spec.get("italic", False),
            center=spec.get("center", False),
        )
    return cur


def main() -> None:
    doc = Document(str(SRC))
    texts = [p.text.strip() for p in doc.paragraphs]
    if any(t == TRACK_HEADING for t in texts):
        print("tracking appendix already present; rewriting file copies only")
    else:
        header = None
        for p in doc.paragraphs:
            if p.text.startswith("Quarter/Year"):
                header = p
                break
        if header is None:
            raise SystemExit("header anchor not found")

        insert_block(
            header,
            [
                {},
                {"text": EPIGRAPH_HEADING, "bold": True},
                {"text": "Work is love made visible.", "italic": True, "center": True},
                {"text": "— Kahlil Gibran, The Prophet (1923)", "center": True},
                {
                    "text": (
                        "Epigraph / front matter only. This line is not a scholarly source for "
                        "method, stakeholder salience, SME size, value-creation stakeholder "
                        "theory, or findings (Gibran, 1923)."
                    )
                },
            ],
        )

        tail = doc.paragraphs[-1]
        insert_block(
            tail,
            [
                {},
                {"text": TRACK_HEADING, "bold": True},
                {
                    "text": (
                        "Use this appendix to track what changed in this Project Plan. "
                        "It is a usage log. It is not literature, method, or findings. "
                        "Marc’s comments remain in the file. Spoken interview stems "
                        "(ITDR-GQI-INT-v0.1.1) are unchanged."
                    )
                },
                {"text": "Locked study facts", "bold": True},
                {
                    "text": (
                        "Program / specialization: PhD in Business Management, Information "
                        "Technology Management. Method: generic qualitative inquiry with "
                        "critical-incident interviews. Sample locked at 12. Collection: "
                        "interviews only. SME screen: 10–200 personnel (SBA, 2026 supplies "
                        "the federal small-firm bound; the study still tightens to 10–200). "
                        "Foundation: Donaldson and Preston (1995); Freeman et al. (2004). "
                        "Needed construct: value-creation stakeholder theory (Freeman, "
                        "Phillips, & Sisodia, 2020) on Need / What We Know / Alignment only. "
                        "Operational lens: Mitchell et al. (1997) power, legitimacy, and urgency."
                    )
                },
                {"text": "What changed in this file", "bold": True},
                {
                    "text": (
                        "Alignment to the Program of Study now names the Information "
                        "Technology Management specialization and states that ITDRPaaS is "
                        "an IT-management governance problem (decision rights, escalation, "
                        "defensible proof), not a technical restore script. Freeman et al. "
                        "(2020) supplies the needed construct. Mitchell et al. (1997) remains "
                        "the operational salience lens."
                    )
                },
                {
                    "text": (
                        "Week 8 closures already in the body: sources on the general problem; "
                        "gap wording (removed “to be GQI experienced”); Term. Definition "
                        "format; sample locked at 12; interviews only; SME-panel and mock/pilot "
                        "as next steps; four-column RQ / interview-question / RQ-alignment / "
                        "framework-alignment matrix; credibility and audit trail."
                    )
                },
                {
                    "text": (
                        "U.S.-preference in-text swaps already applied: SBA (2026); Ampel et al. "
                        "(2024); Morse (2015); Paulus (2023); Guest, Namey, and Chen (2020); "
                        "Resnik (2018); Saldaña (2021); Tracy (2010); Gremler (2004) with "
                        "Flanagan (1954) kept for CIT. Dated non-seminal leftovers remain BOLD. "
                        "A full APA 7th pass of leftover list rows (C34/C35) is still open."
                    )
                },
                {"text": "Gibran statements for inclusion", "bold": True},
                {
                    "text": (
                        "Source: Gibran, K. (1923). The prophet. Alfred A. Knopf. U.S. public "
                        "domain. Include one line as an epigraph unless a committee asks for more. "
                        "Do not use any line as support for SME size, salience, VCST, method, "
                        "ethics, findings, or theme titles. Gibran does not replace Freeman et al. (2020)."
                    )
                },
                {"text": "Recommended primary (On Work) — placed as the front-matter epigraph above.", "bold": True},
                {"text": "Work is love made visible.", "italic": True},
                {
                    "text": (
                        "Place: title page, dedication, or opening of Alignment / Need. "
                        "Not this: method source; a construct; proof that recovery work is “love.” "
                        "Optional companion from the same sermon: “And all work is empty save when there is love.”"
                    )
                },
                {"text": "Recommended secondary (On Giving) — cooperative recovery", "bold": True},
                {
                    "text": (
                        "You give but little when you give of your possessions. It is when you "
                        "give of yourself that you truly give."
                    ),
                    "italic": True,
                },
                {
                    "text": (
                        "Place: front matter near the VCST sentence, or a Chapter I epigraph on "
                        "cooperative recovery. Not this: a substitute for Freeman, Phillips, and Sisodia (2020)."
                    )
                },
                {"text": "Optional (On Laws) — plans versus enacted decision rights", "bold": True},
                {"text": "You delight in laying down laws, yet you delight more in breaking them.", "italic": True},
                {
                    "text": (
                        "Place: epigraph before Constructs or Measures. Not this: NIST / ISO / BCM "
                        "scoring; a finding that managers break policy."
                    )
                },
                {"text": "Optional (On Talking) — interview / CIT stance", "bold": True},
                {"text": "You talk when you cease to be at peace with your thoughts.", "italic": True},
                {
                    "text": (
                        "Place: epigraph before the interview protocol or Chapter III collection. "
                        "Not this: a GQI, CIT, or saturation citation."
                    )
                },
                {"text": "Optional (On Reason and Passion) — Chapter I tone", "bold": True},
                {
                    "text": "Your reason and your passion are the rudder and the sails of your seafaring soul.",
                    "italic": True,
                },
                {
                    "text": (
                        "Place: optional Chapter I tone line. Not this: theory; power-legitimacy-urgency; "
                        "recoverability assurance."
                    )
                },
                {"text": "Still open", "bold": True},
                {
                    "text": (
                        "C34/C35 leftover APA 7th rows remain BOLD. Lester et al. (2020) stays in-text "
                        "until a same-length write-up swap is verified. Official RQ wording in this file "
                        "may replace workspace PQ1/PQ2 labels without changing the unit of analysis. "
                        "Do not silently revert Gremler (2004) to Chell/Butterfield in-text; keep Flanagan (1954)."
                    )
                },
            ],
        )
        print("inserted epigraph and tracking appendix")

    doc.save(str(OUT))
    shutil.copy2(OUT, ART)
    shutil.copy2(OUT, ALIAS)
    shutil.copy2(OUT, ART_ALIAS)
    print("saved", OUT, ART, ALIAS, ART_ALIAS)


if __name__ == "__main__":
    main()
