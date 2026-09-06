#!/usr/bin/env python3
"""Prefer U.S.-based references in MW_Wk5_to_wk8_Edit_6Sep26.docx.

Same-length: in-text and list swaps only; no new paragraphs. Seminal/foundational
theory is kept even when not U.S.-authored. Non-U.S. leftovers without a safe
current U.S. swap stay BOLD. Word comment XML is preserved.
"""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

SRC = Path("/workspace/downloads/MW_Wk5_to_wk8_Edit_6Sep26.docx")
OUT = Path("/workspace/downloads/MW_Wk5_to_wk8_Edit_6Sep26.docx")
ART = Path("/opt/cursor/artifacts/MW_Wk5_to_wk8_Edit_6Sep26.docx")


def set_run_font(run, *, size=12, bold=None, italic=None):
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def replace_text(p, text: str) -> None:
    """Replace visible text; keep extra runs empty so comment XML stays."""
    if not p.runs:
        run = p.add_run(text)
        set_run_font(run)
        return
    p.runs[0].text = text
    if p.runs[0].font.name is None:
        set_run_font(p.runs[0])
    p.runs[0].bold = False
    for r in p.runs[1:]:
        r.text = ""


def clear_hyperlinks(p) -> None:
    for h in list(p._p.findall(qn("w:hyperlink"))):
        p._p.remove(h)


def add_hyperlink(paragraph, text: str, url: str, *, bold=False):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rpr.append(rfonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "24")
    rpr.append(sz)
    if bold:
        b = OxmlElement("w:b")
        rpr.append(b)
    new_run.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def write_reference(p, parts, url=None, *, bold=False):
    clear_hyperlinks(p)
    if not p.runs:
        p.add_run("")
    for r in p.runs:
        r.text = ""
        r.bold = False
        r.italic = False
    first = True
    for text, italic in parts:
        if first and p.runs:
            run = p.runs[0]
            run.text = text
            first = False
        else:
            run = p.add_run(text)
        set_run_font(run, bold=bold, italic=italic)
    if url:
        add_hyperlink(p, url, url, bold=bold)


# Longer strings first so Chell/Butterfield clusters collapse cleanly.
INTEXT_SWAPS = [
    ("Flanagan, 1954; Chell, 2004; Butterfield et al., 2005", "Flanagan, 1954; Gremler, 2004"),
    ("Flanagan, 1954; Chell, 2004", "Flanagan, 1954; Gremler, 2004"),
    ("Ekinci et al., 2025", "U.S. Small Business Administration, 2026"),
    ("Molete et al., 2025", "Ampel et al., 2024"),
    ("Korstjens & Moser, 2018", "Morse, 2015"),
    ("Carcary, 2020", "Paulus, 2023"),
    ("Naeem et al., 2024", "Guest et al., 2020"),
    ("Miteu, 2024", "Resnik, 2018"),
    ("Braun & Clarke, 2019", "Saldaña, 2021"),
    ("Nowell et al., 2017", "Tracy, 2010"),
]


# bold=True = not U.S.-based and not seminal/foundational (left in list).
REFS = [
    (
        [
            ("Agle, B. R., Mitchell, R. K., & Sonnenfeld, J. A. (1999). Who matters to CEOs? An investigation of stakeholder attributes and salience, corporate performance, and CEO values. ", False),
            ("Academy of Management Journal, 42", True),
            ("(5), 507–525. ", False),
        ],
        "https://doi.org/10.5465/256973",
        False,
    ),
    (
        [
            ("Agle, B. R., Donaldson, T., Freeman, R. E., Jensen, M. C., Mitchell, R. K., & Wood, D. J. (2008). Dialogue: Toward superior stakeholder theory. ", False),
            ("Business Ethics Quarterly, 18", True),
            ("(2), 153–190. ", False),
        ],
        "https://doi.org/10.5840/beq200818214",
        False,
    ),
    (
        [
            ("Ampel, B. M., Samtani, S., Zhu, H., Chen, H., & Nunamaker, J. F., Jr. (2024). Improving threat mitigation through a cybersecurity risk management framework: A computational design science approach. ", False),
            ("Journal of Management Information Systems, 41", True),
            ("(1), 236–265. ", False),
        ],
        "https://doi.org/10.1080/07421222.2023.2301178",
        False,
    ),
    (
        [
            ("Brinkmann, S., & Kvale, S. (2017). Ethics in qualitative psychological research. In C. Willig & W. Stainton Rogers (Eds.), ", False),
            ("The SAGE handbook of qualitative research in psychology", True),
            (" (2nd ed., pp. 259–273). SAGE. ", False),
        ],
        None,
        True,
    ),
    (
        [
            ("Butterfield, L. D., Borgen, W. A., Amundson, N. E., & Maglio, A.-S. T. (2005). Fifty years of the critical incident technique: 1954–2004 and beyond. ", False),
            ("Qualitative Research, 5", True),
            ("(4), 475–497. ", False),
        ],
        "https://doi.org/10.1177/1468794105056924",
        True,
    ),
    (
        [
            ("Caelli, K., Ray, L., & Mill, J. (2003). “Clear as mud”: Toward greater clarity in generic qualitative research. ", False),
            ("International Journal of Qualitative Methods, 2", True),
            ("(2), 1–13. ", False),
        ],
        "https://doi.org/10.1177/160940690300200201",
        False,
    ),
    (
        [
            ("Capella University. (2024). ", False),
            ("University catalog: Doctor of Philosophy (PhD) in Business Management, Information Technology Management specialization", True),
            (". Capella University.", False),
        ],
        None,
        False,
    ),
    (
        [
            ("Cooper, S., & Endacott, R. (2007). Generic qualitative research: A design for qualitative research in emergency care? ", False),
            ("Emergency Medicine Journal, 24", True),
            ("(12), 816–819. ", False),
        ],
        "https://doi.org/10.1136/emj.2007.050641",
        True,
    ),
    (
        [
            ("Creswell, J. W., & Miller, D. L. (2000). Determining validity in qualitative inquiry. ", False),
            ("Theory Into Practice, 39", True),
            ("(3), 124–130. ", False),
        ],
        "https://doi.org/10.1207/s15430421tip3903_2",
        False,
    ),
    (
        [
            ("Creswell, J. W., & Poth, C. N. (2024). ", False),
            ("Qualitative inquiry and research design: Choosing among five approaches", True),
            (" (5th ed.). SAGE.", False),
        ],
        None,
        False,
    ),
    (
        [
            ("Donaldson, T., & Preston, L. E. (1995). The stakeholder theory of the corporation: Concepts, evidence, and implications. ", False),
            ("Academy of Management Review, 20", True),
            ("(1), 65–91. ", False),
        ],
        "https://doi.org/10.5465/amr.1995.9503271992",
        False,
    ),
    (
        [
            ("Dorobantu, S., Henisz, W. J., & Nartey, L. J. (2024). Firm–stakeholder dialogue and the media: The evolution of stakeholder evaluations in different informational environments. ", False),
            ("Academy of Management Journal, 67", True),
            ("(1), 92–125. ", False),
        ],
        "https://doi.org/10.5465/amj.2021.0103",
        False,
    ),
    (
        [
            ("Erdiaw-Kwasie, M. O., Alam, K., & Shahiduzzaman, M. (2017). Towards understanding stakeholder salience transition and relational approach to ‘better’ corporate social responsibility: A case for a proposed model in practice. ", False),
            ("Journal of Business Ethics, 144", True),
            ("(1), 85–101. ", False),
        ],
        "https://doi.org/10.1007/s10551-015-2805-z",
        True,
    ),
    (
        [
            ("Farahbod, K., Shayo, C., & Varzandeh, J. (2022). Six sigma and lean operations in cybersecurity management. ", False),
            ("Journal of Business and Behavioral Sciences, 34", True),
            ("(1), 99–109.", False),
        ],
        None,
        False,
    ),
    (
        [
            ("Federal Bureau of Investigation. (2024). ", False),
            ("Internet crime report 2024", True),
            (". ", False),
        ],
        "https://www.ic3.gov/AnnualReport/Reports/2024_IC3Report.pdf",
        False,
    ),
    (
        [
            ("Fereday, J., & Muir-Cochrane, E. (2006). Demonstrating rigor using thematic analysis: A hybrid approach of inductive and deductive coding and theme development. ", False),
            ("International Journal of Qualitative Methods, 5", True),
            ("(1), 80–92. ", False),
        ],
        "https://doi.org/10.1177/160940690600500107",
        False,
    ),
    (
        [
            ("Fischer, D., Brettel, M., & Mauer, R. (2020). The three dimensions of sustainability: A delicate balancing act for entrepreneurs made more complex by stakeholder expectations. ", False),
            ("Journal of Business Ethics, 163", True),
            ("(1), 87–106. ", False),
        ],
        "https://doi.org/10.1007/s10551-018-4012-1",
        True,
    ),
    (
        [
            ("Flanagan, J. C. (1954). The critical incident technique. ", False),
            ("Psychological Bulletin, 51", True),
            ("(4), 327–358. ", False),
        ],
        "https://doi.org/10.1037/h0061470",
        False,
    ),
    (
        [
            ("Freeman, R. E. (1984). ", False),
            ("Strategic management: A stakeholder approach", True),
            (". Pitman.", False),
        ],
        None,
        False,
    ),
    (
        [
            ("Freeman, R. E. (2008). ", False),
            ("Managing for stakeholders: Survival, reputation, and success", True),
            (". Yale University Press.", False),
        ],
        None,
        False,
    ),
    (
        [
            ("Freeman, R. E. (2009). Stakeholder theory: 25 years later. ", False),
            ("Philosophy of Management, 8", True),
            ("(3), 97–107. ", False),
        ],
        "https://doi.org/10.5840/pom20098310",
        False,
    ),
    (
        [
            ("Freeman, R. E., Gilbert, D. R., Jr., & Hartman, E. (1988). Values and the foundations of strategic management. ", False),
            ("Journal of Business Ethics, 7", True),
            ("(11), 821–834. ", False),
        ],
        "https://doi.org/10.1007/BF00383045",
        False,
    ),
    (
        [
            ("Freeman, R. E., Martin, K., & Parmar, B. (2007). Stakeholder capitalism. ", False),
            ("Journal of Business Ethics, 74", True),
            ("(4), 303–314. ", False),
        ],
        "https://doi.org/10.1007/s10551-007-9517-y",
        False,
    ),
    (
        [
            ("Freeman, R. E., Phillips, R., & Sisodia, R. (2020). Tensions in stakeholder theory. ", False),
            ("Business & Society, 59", True),
            ("(2), 213–231. ", False),
        ],
        "https://doi.org/10.1177/0007650318773750",
        False,
    ),
    (
        [
            ("Freeman, R. E., Wicks, A. C., & Parmar, B. (2004). Stakeholder theory and the corporate objective revisited. ", False),
            ("Organization Science, 15", True),
            ("(3), 364–369. ", False),
        ],
        "https://doi.org/10.1287/orsc.1040.0066",
        False,
    ),
    (
        [
            ("Fusch, P., Fusch, G. E., & Ness, L. R. (2018). Denzin’s paradigm shift: Revisiting triangulation in qualitative research. ", False),
            ("Journal of Social Change, 10", True),
            ("(1), 19–32. ", False),
        ],
        "https://doi.org/10.5590/JOSC.2018.10.1.02",
        False,
    ),
    (
        [
            ("Gabriel, C., & Shafique, K. (2024). An ethical salience framework to achieve sustainable development goals. ", False),
            ("Sustainable Development, 32", True),
            ("(4), 3213–3225. ", False),
        ],
        "https://doi.org/10.1002/sd.2840",
        True,
    ),
    (
        [
            ("Gremler, D. D. (2004). The critical incident technique in service research. ", False),
            ("Journal of Service Research, 7", True),
            ("(1), 65–89. ", False),
        ],
        "https://doi.org/10.1177/1094670504266138",
        False,
    ),
    (
        [
            ("Guba, E. G., & Lincoln, Y. S. (1994). Competing paradigms in qualitative research. In N. K. Denzin & Y. S. Lincoln (Eds.), ", False),
            ("Handbook of qualitative research", True),
            (" (pp. 105–117). SAGE.", False),
        ],
        None,
        False,
    ),
    (
        [
            ("Guest, G., Bunce, A., & Johnson, L. (2006). How many interviews are enough? An experiment with data saturation and variability. ", False),
            ("Field Methods, 18", True),
            ("(1), 59–82. ", False),
        ],
        "https://doi.org/10.1177/1525822X05279903",
        False,
    ),
    (
        [
            ("Guest, G., Namey, E., & Chen, M. (2020). A simple method to assess and report thematic saturation in qualitative research. ", False),
            ("PLOS ONE, 15", True),
            ("(5), Article e0232076. ", False),
        ],
        "https://doi.org/10.1371/journal.pone.0232076",
        False,
    ),
    (
        [
            ("Kahlke, R. M. (2014). Generic qualitative approaches: Pitfalls and benefits of methodological mixology. ", False),
            ("International Journal of Qualitative Methods, 13", True),
            ("(1), 37–52. ", False),
        ],
        "https://doi.org/10.1177/160940691401300119",
        False,
    ),
    (
        [
            ("Kenny, C., Weber, C., & Bratton, K. (2017). The characteristics of interpersonal networks in disaster response. ", False),
            ("Social Science Quarterly, 98", True),
            ("(2), 566–583. ", False),
        ],
        "https://doi.org/10.1111/ssqu.12328",
        False,
    ),
    (
        [
            ("Koerber, A., & McMichael, L. (2008). Qualitative sampling methods: A primer for technical communicators. ", False),
            ("Journal of Business and Technical Communication, 22", True),
            ("(4), 454–473. ", False),
        ],
        "https://doi.org/10.1177/1050651908320362",
        False,
    ),
    (
        [
            ("Kvale, S., & Brinkmann, S. (2015). ", False),
            ("InterViews: Learning the craft of qualitative research interviewing", True),
            (" (3rd ed.). SAGE.", False),
        ],
        None,
        True,
    ),
    (
        [
            ("Lester, J. N., Cho, Y., & Lochmiller, C. R. (2020). Learning to do qualitative data analysis: A starting point. ", False),
            ("Human Resource Development Review, 19", True),
            ("(1), 94–106. ", False),
        ],
        "https://doi.org/10.1177/1534484320903890",
        False,
    ),
    (
        [
            ("Lincoln, Y. S., & Guba, E. G. (1985). ", False),
            ("Naturalistic inquiry", True),
            (". SAGE.", False),
        ],
        None,
        False,
    ),
    (
        [
            ("Liu, W., Liu, R., Chen, H., & Mboga, J. (2020). Perspectives on disruptive technology and innovation. ", False),
            ("International Journal of Conflict Management, 31", True),
            ("(3), 313–331. ", False),
        ],
        "https://doi.org/10.1108/IJCMA-09-2019-0172",
        True,
    ),
    (
        [
            ("Lowry, M. R., Vance, A., & Vance, M. D. (2025). Inexpert supervision: Field evidence on boards’ oversight of cybersecurity. ", False),
            ("Management Science", True),
            (". ", False),
        ],
        "https://doi.org/10.1287/mnsc.2023.04147",
        False,
    ),
    (
        [
            ("Malterud, K., Siersma, V. D., & Guassora, A. D. (2016). Sample size in qualitative interview studies: Guided by information power. ", False),
            ("Qualitative Health Research, 26", True),
            ("(13), 1753–1760. ", False),
        ],
        "https://doi.org/10.1177/1049732315617444",
        False,
    ),
    (
        [
            ("Mitchell, R. K., Agle, B. R., & Wood, D. J. (1997). Toward a theory of stakeholder identification and salience: Defining the principle of who and what really counts. ", False),
            ("Academy of Management Review, 22", True),
            ("(4), 853–886. ", False),
        ],
        "https://doi.org/10.5465/amr.1997.9711022105",
        False,
    ),
    (
        [
            ("Mojtahedi, M., & Oo, B. L. (2017). Critical attributes for proactive engagement of stakeholders in disaster risk management. ", False),
            ("International Journal of Disaster Risk Reduction, 21", True),
            (", 35–43. ", False),
        ],
        "https://doi.org/10.1016/j.ijdrr.2016.10.017",
        True,
    ),
    (
        [
            ("Molete, O. B., Mokhele, S. E., Ntombela, S. D., & Thango, B. A. (2025). The impact of IT strategic planning process on SME performance: A systematic review. ", False),
            ("Businesses, 5", True),
            ("(1), Article 2. ", False),
        ],
        "https://doi.org/10.3390/businesses5010002",
        True,
    ),
    (
        [
            ("Morse, J. M. (2015). Critical analysis of strategies for determining rigor in qualitative inquiry. ", False),
            ("Qualitative Health Research, 25", True),
            ("(9), 1212–1222. ", False),
        ],
        "https://doi.org/10.1177/1049732315588501",
        False,
    ),
    (
        [
            ("Narayanan, A., Finucane, M., Acosta, J., & Wicker, A. (2020). From awareness to action: Accounting for infrastructure interdependencies in disaster response and recovery planning. ", False),
            ("GeoHealth, 4", True),
            ("(7), Article e2020GH000251. ", False),
        ],
        "https://doi.org/10.1029/2020GH000251",
        False,
    ),
    (
        [
            ("National Commission for the Protection of Human Subjects of Biomedical and Behavioral Research. (1979). ", False),
            ("The Belmont report: Ethical principles and guidelines for the protection of human subjects of research", True),
            (". U.S. Department of Health, Education, and Welfare. ", False),
        ],
        "https://www.hhs.gov/ohrp/regulations-and-policy/belmont-report/index.html",
        False,
    ),
    (
        [
            ("Nowell, L. S., Norris, J. M., White, D. E., & Moules, N. J. (2017). Thematic analysis: Striving to meet the trustworthiness criteria. ", False),
            ("International Journal of Qualitative Methods, 16", True),
            ("(1), 1–13. ", False),
        ],
        "https://doi.org/10.1177/1609406917733847",
        True,
    ),
    (
        [
            ("Nowell, B., & Albrecht, K. (2019). A reviewer’s guide to qualitative rigor. ", False),
            ("Journal of Public Administration Research and Theory, 29", True),
            ("(2), 348–363. ", False),
        ],
        "https://doi.org/10.1093/jopart/muy052",
        False,
    ),
    (
        [
            ("Örtenblad, A. (2007). Senge’s many faces: Problem or opportunity? ", False),
            ("The Learning Organization, 14", True),
            ("(2), 108–122. ", False),
        ],
        "https://doi.org/10.1108/09696470710726989",
        True,
    ),
    (
        [
            ("Park, J., Son, Y., & Angst, C. M. (2023). The value of centralized IT in building resilience during crises: Evidence from U.S. higher education’s transition to emergency remote teaching. ", False),
            ("MIS Quarterly, 47", True),
            ("(1), 451–482. ", False),
        ],
        "https://doi.org/10.25300/MISQ/2022/17265",
        False,
    ),
    (
        [
            ("Paulus, T. M. (2023). Using qualitative data analysis software to support digital research workflows. ", False),
            ("Human Resource Development Review, 22", True),
            ("(1), 139–148. ", False),
        ],
        "https://doi.org/10.1177/15344843221138381",
        False,
    ),
    (
        [
            ("Percy, W. H., Kostere, K., & Kostere, S. (2015). Generic qualitative research in psychology. ", False),
            ("The Qualitative Report, 20", True),
            ("(2), 76–85. ", False),
        ],
        "https://doi.org/10.46743/2160-3715/2015.2097",
        False,
    ),
    (
        [
            ("Resnik, D. B. (2018). ", False),
            ("The ethics of research with human subjects: Protecting people, advancing science, promoting trust", True),
            (". Springer. ", False),
        ],
        "https://doi.org/10.1007/978-3-319-68756-8",
        False,
    ),
    (
        [
            ("Saldaña, J. (2021). ", False),
            ("The coding manual for qualitative researchers", True),
            (" (4th ed.). SAGE.", False),
        ],
        None,
        False,
    ),
    (
        [
            ("Stieb, J. A. (2009). Assessing Freeman’s stakeholder theory. ", False),
            ("Journal of Business Ethics, 87", True),
            ("(3), 401–414. ", False),
        ],
        "https://doi.org/10.1007/s10551-008-9928-4",
        False,
    ),
    (
        [
            ("Taylor, S. J., & Bogdan, R. (1998). ", False),
            ("Introduction to qualitative research methods: A guidebook and resource", True),
            (" (3rd ed.). Wiley.", False),
        ],
        None,
        False,
    ),
    (
        [
            ("Tisdell, E. J., Merriam, S. B., & Stuckey-Peyrot, H. L. (2025). ", False),
            ("Qualitative research: A guide to design and implementation", True),
            (". John Wiley & Sons.", False),
        ],
        None,
        False,
    ),
    (
        [
            ("Tracy, S. J. (2010). Qualitative quality: Eight “big-tent” criteria for excellent qualitative research. ", False),
            ("Qualitative Inquiry, 16", True),
            ("(10), 837–851. ", False),
        ],
        "https://doi.org/10.1177/1077800410383121",
        False,
    ),
    (
        [
            ("U.S. Small Business Administration. (2026). ", False),
            ("Frequently asked questions about small business", True),
            (". Office of Advocacy. ", False),
        ],
        "https://advocacy.sba.gov/wp-content/uploads/2026/02/FINAL_FAQsAboutSmallBusiness_2026_012826.pdf",
        False,
    ),
    (
        [
            ("van der Kamp, M., Tjemkes, B., Duplat, V., & Jehn, K. (2023). On alliance teams: Conceptualization, review, and future research agenda. ", False),
            ("Human Relations, 76", True),
            ("(9), 1382–1413. ", False),
        ],
        "https://doi.org/10.1177/00187267221104985",
        True,
    ),
    (
        [
            ("de Vaujany, F.-X., Leclercq-Vandelannoitte, A., Aroles, J., Introna, L., & Davidson, S. (2025). Rethinking responsibility in the digital age: A narrative approach. ", False),
            ("MIS Quarterly, 49", True),
            ("(4), 1295–1318.", False),
        ],
        None,
        True,
    ),
    (
        [
            ("Verizon. (2025). ", False),
            ("2025 data breach investigations report", True),
            (". ", False),
        ],
        "https://www.verizon.com/business/resources/reports/dbir/",
        False,
    ),
    (
        [
            ("Vos, J. (2022). A delicate balance: How structure, creativity and innovation play a role in advancing change. ", False),
            ("HCM Sales, Marketing & Alliance Excellence, 21", True),
            ("(11), 15–17.", False),
        ],
        None,
        True,
    ),
    (
        [
            ("Walker, M. (2026). ", False),
            ("GQI semi-structured interview guide: Where leverage sits", True),
            (" (Instrument ID ITDR-GQI-INT-v0.1.1) [Unpublished research instrument]. Capella University.", False),
        ],
        None,
        False,
    ),
    (
        [
            ("Weick, K. E. (1989). Theory construction as disciplined imagination. ", False),
            ("Academy of Management Review, 14", True),
            ("(4), 516–531. ", False),
        ],
        "https://doi.org/10.5465/amr.1989.4308376",
        False,
    ),
    (
        [
            ("Weill, P., & Ross, J. W. (2004). ", False),
            ("IT governance: How top performers manage IT decision rights for superior results", True),
            (". Harvard Business School Press.", False),
        ],
        None,
        False,
    ),
    (
        [
            ("Williams, T. A., Gruber, D. A., Sutcliffe, K. M., Shepherd, D. A., & Zhao, E. Y. (2017). Organizational response to adversity: Fusing crisis management and resilience research streams. ", False),
            ("Academy of Management Annals, 11", True),
            ("(2), 733–769. ", False),
        ],
        "https://doi.org/10.5465/annals.2015.0134",
        False,
    ),
    (
        [
            ("Wutich, A., Beresford, M., & Bernard, H. R. (2024). Sample sizes for 10 types of qualitative data analysis: An integrative review, empirical guidance, and next steps. ", False),
            ("International Journal of Qualitative Methods, 23", True),
            (", 1–23. ", False),
        ],
        "https://doi.org/10.1177/16094069241296206",
        False,
    ),
]


def patch_body(paragraphs) -> int:
    n = 0
    for p in paragraphs[:154]:
        t = p.text
        new = t
        for old, repl in INTEXT_SWAPS:
            if old in new:
                new = new.replace(old, repl)
        if new != t:
            replace_text(p, new)
            n += 1
    return n


def main():
    if len(REFS) != 68:
        raise SystemExit(f"expected 68 refs to fill existing slots, got {len(REFS)}")

    doc = Document(SRC)
    p = doc.paragraphs
    if len(p) != 228:
        raise SystemExit(f"unexpected paragraph count {len(p)}")

    swapped_paras = patch_body(p)

    first_ref = 157
    last_existing = 227
    for i, (parts, url, bold) in enumerate(REFS):
        idx = first_ref + i
        if idx > last_existing:
            raise SystemExit(f"need extra paragraph at {idx}")
        write_reference(p[idx], parts, url, bold=bold)

    for idx in range(first_ref + len(REFS), last_existing + 1):
        clear_hyperlinks(p[idx])
        replace_text(p[idx], "")

    if len(doc.paragraphs) != 228:
        raise SystemExit(f"paragraph count drifted to {len(doc.paragraphs)}")

    doc.save(OUT)
    ART.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUT, ART)
    print("wrote", OUT)
    print("body paragraphs patched", swapped_paras)
    print("refs written", len(REFS))
    print("bold leftover refs", sum(1 for *_, b in REFS if b))


if __name__ == "__main__":
    main()
