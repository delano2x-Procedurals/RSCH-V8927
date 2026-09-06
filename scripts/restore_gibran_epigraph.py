#!/usr/bin/env python3
"""Restore Gibran (1923) as an include-epigraph reference.

Inserts the APA line alphabetically between Gabriel and Gremler.
Does not add body/Alignment paragraphs. Preserves Word comments.
"""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

SRC = Path("/workspace/downloads/MW_Wk5_to_wk8_Edit_6Sep26.docx")
OUT = Path("/workspace/downloads/MW_Wk5_to_wk8_Edit_6Sep26.docx")
ART = Path("/opt/cursor/artifacts/MW_Wk5_to_wk8_Edit_6Sep26.docx")

APA = "Gibran, K. (1923). The prophet. Alfred A. Knopf."


def set_run_font(run, *, size=12, italic=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if italic is not None:
        run.italic = italic


def insert_paragraph_after(paragraph, text: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    # Book title italicized: "The prophet"
    prefix, rest = text.split("The prophet.", 1)
    r1 = new_para.add_run(prefix)
    set_run_font(r1)
    r2 = new_para.add_run("The prophet.")
    set_run_font(r2, italic=True)
    r3 = new_para.add_run(rest)
    set_run_font(r3)
    pf = new_para.paragraph_format
    pf.line_spacing = 2.0
    pf.left_indent = Inches(0.5)
    pf.first_line_indent = Inches(-0.5)
    return new_para


def main() -> None:
    doc = Document(str(SRC))
    existing = [p for p in doc.paragraphs if p.text.startswith("Gibran, K.")]
    if existing:
        print("Gibran already present; no insert")
    else:
        anchor = None
        for p in doc.paragraphs:
            if p.text.startswith("Gabriel, C."):
                anchor = p
                break
        if anchor is None:
            raise SystemExit("Gabriel anchor not found")
        insert_paragraph_after(anchor, APA)
        print("inserted Gibran after Gabriel")
    doc.save(str(OUT))
    shutil.copy2(OUT, ART)
    print("saved", OUT, "and", ART)


if __name__ == "__main__":
    main()
