#!/usr/bin/env python3
"""Scholarly-clarity edits + citation alignment with bolded changes.

Preserves Word comments by clearing run text and rewriting runs in place
(plus extra runs as needed). Does not delete comment-range XML.
"""

from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

SRC = Path("/workspace/downloads/MW_Wk5_to_wk8_Edit_6Sep26.docx")
OUT = Path("/workspace/downloads/MW_Wk5_to_wk8_Edit_6Sep26.docx")
ART = Path("/opt/cursor/artifacts/MW_Wk5_to_wk8_Edit_6Sep26.docx")
ALIAS = Path("/workspace/downloads/MW_Wk5_to_wk8_Edit_6Sep26_updated.docx")
ART_ALIAS = Path("/opt/cursor/artifacts/MW_Wk5_to_wk8_Edit_6Sep26_updated.docx")


def set_run_font(run, *, size=12, bold=None, italic=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def write_marked(p, marked: str) -> None:
    """Write paragraph text. Spans wrapped in ** ** are bold (updated citations)."""
    parts = re.split(r"\*\*(.+?)\*\*", marked)
    segs: list[tuple[str, bool]] = []
    for i, part in enumerate(parts):
        if not part:
            continue
        segs.append((part, i % 2 == 1))
    if not segs:
        return
    if not p.runs:
        for text, bold in segs:
            r = p.add_run(text)
            set_run_font(r, bold=bold)
        return
    p.runs[0].text = segs[0][0]
    set_run_font(p.runs[0], bold=segs[0][1])
    for r in p.runs[1:]:
        r.text = ""
    for text, bold in segs[1:]:
        r = p.add_run(text)
        set_run_font(r, bold=bold)


def find_para(doc, prefix: str):
    for p in doc.paragraphs:
        if p.text.startswith(prefix):
            return p
    raise SystemExit(f"anchor not found: {prefix[:60]}")


def main() -> None:
    doc = Document(str(SRC))

    write_marked(
        find_para(doc, "This study adopts a constructivist"),
        "This study adopts a constructivist–interpretive worldview: organizational "
        "realities are constructed through participants’ interactions and interpretations "
        "rather than existing as a single objective reality "
        "(**Creswell & Poth, 2024**; Guba & Lincoln, 1994). Aligning with a "
        "subjectivist–transactional epistemology, knowledge of ITDRPaaS governance "
        "emerges from the interaction between the researcher and participants "
        "(Guba & Lincoln, 1994). That stance supports generic qualitative inquiry "
        "aimed at interpreting how managers understand recovery priorities, decision "
        "authority, escalation requirements, and acceptable evidence of recoverability. "
        "The design therefore describes how IT and network managers, read through "
        "stakeholder-salience, translate organizational IT strategy into operational "
        "disaster-recovery actions rather than testing predefined compliance checklists. "
        "In SME environments, competing stakeholder demands routinely complicate recovery "
        "prioritization, restoration authority, and acceptable service-recovery thresholds. "
        "Accordingly, the design captures how those demands condition who counts in "
        "recovery prioritization, who decides and escalates under pressure, and what "
        "counts as an acceptable recovery standard.",
    )

    write_marked(
        find_para(doc, "Stakeholder-salience theory will serve"),
        "Stakeholder-salience theory is the interpretive foundation for understanding "
        "how perceptions of power, legitimacy, and urgency influence the prioritization "
        "of recovery actions during time-critical events (Mitchell et al., 1997). Within "
        "this generic qualitative inquiry, the phenomenon is IT and network managers’ "
        "accounts of governance and decision dynamics as they interpret stakeholder "
        "influence during disruption or recovery planning. Managers must convert IT "
        "strategy into operational recovery authority while navigating conflicting claims "
        "about what should be restored first and who may authorize action. The approach "
        "yields descriptive insight into decision rights and escalation pathways without "
        "reducing the inquiry to phenomenology (Weill & Ross, 2004). It also examines "
        "evidentiary standards that support recoverability assurance when managers must "
        "defend recovery outcomes to stakeholders (Lowry et al., 2025; Park et al., 2023). "
        "Power, legitimacy, and urgency therefore keep the method aligned with the study’s purpose.",
    )

    write_marked(
        find_para(doc, "The target population comprises IT and network managers"),
        "The target population comprises IT and network managers (or equivalent titles "
        "such as systems, infrastructure, or IT operations managers) employed by U.S. "
        "small and medium enterprises of 10 to 200 employees "
        "(**U.S. Small Business Administration, 2026**). Eligible participants have "
        "responsibility for, or direct participation in, IT disaster recovery involving "
        "an external platform or managed recovery service (ITDRPaaS or functionally "
        "equivalent cloud recovery). Population characteristics therefore center on role "
        "authority, U.S. SME employment, and recent platform-mediated recovery exposure "
        "rather than executive-only or vendor-side perspectives. The study is not limited "
        "to a single organizational site; participants will be drawn from multiple U.S. "
        "SME contexts in which IT systems are operationally critical. The locked sample "
        "is 12 interviews. That number is justified by information power for a narrow aim "
        "and a specific population (Malterud et al., 2016) and is the school-required "
        "fixed count rather than a range. Guest et al. (2006) is reserved for later "
        "stopping discussion only and is not the interview method.",
    )

    write_marked(
        find_para(doc, "Operational procedures begin with purposive"),
        "Operational procedures begin with purposive and criterion sampling to identify "
        "qualified participants (Percy et al., 2015). The locked sample is 12. "
        "Inclusion. The participant serves as an IT or network manager (or equivalent "
        "systems, infrastructure, or IT operations role) in a U.S. SME of 10 to 200 "
        "personnel (**U.S. Small Business Administration, 2026**); holds direct "
        "operational responsibility for, or participation in, IT disaster-recovery "
        "coordination that uses an external platform or managed recovery service; and "
        "can name at least one disruption, failover, recovery test, or operational "
        "recovery in the preceding 36 months. Exclusion. Vendor-only employees and "
        "managers who cannot name a qualifying event. After eligibility is verified, "
        "the researcher emails the informed consent form; once signed electronic consent "
        "is returned, a virtual, recorded English interview is scheduled.",
    )

    write_marked(
        find_para(doc, "This study will use a non-probability"),
        "This study will use a non-probability, criterion-based purposive sampling "
        "strategy to select information-rich participants whose professional roles and "
        "recent recovery experience align with the research purpose (Percy et al., 2015; "
        "**Creswell & Poth, 2024**). The target population consists of U.S. SME IT, "
        "network, systems, infrastructure, or operations managers in firms of 10 to 200 "
        "personnel who have direct responsibility for, or participation in, IT disaster "
        "recovery that uses an external recovery platform, managed recovery service, "
        "ITDRPaaS, or a functionally equivalent cloud-recovery capability. Inclusion. "
        "At least one disruption, failover, recovery test, or operational recovery in "
        "the preceding 36 months in which that external capability was used. Exclusion. "
        "Individuals employed only by a vendor, or who lack direct platform-mediated "
        "recovery responsibility. Criterion-based purposive sampling is appropriate for "
        "this generic qualitative inquiry because the researcher will deliberately select "
        "participants who can recount a named recovery event rather than attempt to "
        "obtain a statistically representative sample (Percy et al., 2015).",
    )

    write_marked(
        find_para(doc, "The sample is locked at 12 participants. Sample adequacy"),
        "The sample is locked at 12 participants. Sample adequacy will not be determined "
        "through statistical power calculations (**Wutich et al., 2024**). It will be "
        "justified by information power: a narrow study aim, a highly specific participant "
        "population, theoretically informed but not closed-ended dialogue, and an in-depth "
        "thematic analysis strategy (Malterud et al., 2016). Twelve is the planned "
        "recruitment target for the detailed screening and interview process. It is not "
        "a finding that the gap is closed, nor is there a claim that saturation occurs "
        "at the twelfth interview.",
    )

    write_marked(
        find_para(doc, "Recruitment and preliminary analysis will proceed"),
        "Recruitment and preliminary analysis will proceed through the locked sample of 12. "
        "Stopping will be documented as an analytic decision: the researcher will record "
        "whether interviews in that corpus continue to contribute new meaning units, "
        "categories, or thematic variation relevant to the research questions "
        "(**Guest et al., 2020**). Guest et al. (2006) is used only as background that "
        "some homogeneous interview studies have observed early code stability; that "
        "source will not be used to claim in advance that 12 interviews suffice. The "
        "final participant count and the rationale for ending recruitment will be "
        "recorded in the study audit trail (**Paulus, 2023**).",
    )

    write_marked(
        find_para(doc, "This population definition is the pool"),
        "This population definition is the pool from which the locked sample of 12 "
        "participants will be drawn. It is not itself the sample, and it is not a "
        "statistically representative frame of all U.S. SME recovery settings "
        "(**Creswell & Poth, 2024**; Percy et al., 2015). The definition is written this "
        "tightly because information power depends on sample specificity: a narrow class "
        "of SME recovery events and a dense participant profile increase the information "
        "available from each interview (Malterud et al., 2016). The 12-participant sample, "
        "the inclusion and exclusion screens, and the individual recruitment path all "
        "depend on this population boundary remaining stable.",
    )

    write_marked(
        find_para(doc, "Honesty and integrity will govern"),
        "Honesty and integrity will govern collection, storage, analysis, reporting, and "
        "any later discussion of the study (**Resnik, 2018**). Recordings and RAW notes "
        "will be stored in an encrypted, password-protected directory separate from CLEAN "
        "transcripts and the Delve audit nest. The researcher will not continue a session "
        "from memory if recording fails, will not invent artifact contents a manager did "
        "not state, and will not claim that Delve performed the analysis. Findings will "
        "be disseminated as de-identified patterned accounts of named U.S. SME recovery "
        "events, without employer, owner, vendor, or product names, so the study can "
        "report how managers enacted whose claim counted, who decided, and what counted "
        "as enough proof without exposing the participant or the firm.",
    )

    doc.save(str(OUT))
    shutil.copy2(OUT, ART)
    shutil.copy2(OUT, ALIAS)
    shutil.copy2(OUT, ART_ALIAS)
    print("saved clarity + bold citation updates")


if __name__ == "__main__":
    main()
